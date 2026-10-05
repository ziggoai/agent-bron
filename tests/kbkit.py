"""Knowledge-base test helpers: stand-in models and sample documents. Nothing here calls a real model."""

import ctypes
import ctypes.util
import os
from pathlib import Path

ATTR = "com.google.drivefs.item-id#S"


def fake_drive(tmp_path: Path) -> Path:
    root = tmp_path / "GoogleDrive-test" / "My Drive"
    root.mkdir(parents=True)
    return root


def set_drive_id(path: Path, item_id: str) -> None:
    if hasattr(os, "setxattr"):
        os.setxattr(str(path), ATTR, item_id.encode())
        return
    # Python has no os.setxattr on macOS; call the system library directly.
    libc = ctypes.CDLL(ctypes.util.find_library("c"), use_errno=True)
    data = item_id.encode()
    if libc.setxattr(str(path).encode(), ATTR.encode(), data, len(data), 0, 0) != 0:
        err = ctypes.get_errno()
        raise OSError(err, os.strerror(err), str(path))


# ---- sample documents (written into tmp_path; never real data) ----

def make_text_pdf(path: Path, pages: list[str], full_page_image: bool = False) -> Path:
    """A minimal hand-built PDF: one Helvetica text page per string (ASCII and Latin-1).

    With full_page_image, every page also carries a grey image covering the whole page (a scan with a text header).
    """
    objs: list[bytes] = []
    n = len(pages)
    kids = " ".join(f"{3 + 2 * i} 0 R" for i in range(n))
    objs.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objs.append(f"<< /Type /Pages /Kids [{kids}] /Count {n} >>".encode())
    font_id = 3 + 2 * n
    for i, text in enumerate(pages):
        content_id = 4 + 2 * i
        xobject = f" /XObject << /Im1 {font_id + 1} 0 R >>" if full_page_image else ""
        objs.append(
            (f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
             f"/Resources << /Font << /F1 {font_id} 0 R >>{xobject} >> /Contents {content_id} 0 R >>").encode()
        )
        escaped = text.encode("latin-1", "replace").replace(b"\\", b"\\\\").replace(b"(", b"\\(").replace(b")", b"\\)")
        stream = (b"q 612 0 0 792 0 0 cm /Im1 Do Q " if full_page_image else b"") + b"BT /F1 11 Tf 50 750 Td (" + escaped + b") Tj ET"
        objs.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
    objs.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>")
    if full_page_image:
        pixels = bytes([200]) * 16
        objs.append(b"<< /Type /XObject /Subtype /Image /Width 4 /Height 4 /ColorSpace /DeviceGray "
                    b"/BitsPerComponent 8 /Length 16 >>\nstream\n" + pixels + b"\nendstream")
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objs, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % i + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)
    for off in offsets:
        out += b"%010d 00000 n \n" % off
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, xref)
    path.write_bytes(bytes(out))
    return path


def make_png(path: Path, lines: list[str]) -> Path:
    from PIL import Image, ImageDraw, ImageFont

    font = None
    for name in ("Arial.ttf", "/System/Library/Fonts/Supplemental/Arial.ttf", "DejaVuSans.ttf"):
        try:
            font = ImageFont.truetype(name, 40)
            break
        except OSError:
            continue
    font = font or ImageFont.load_default()
    img = Image.new("RGB", (1400, 100 * max(1, len(lines)) + 60), "white")
    draw = ImageDraw.Draw(img)
    for i, line in enumerate(lines):
        draw.text((40, 40 + 100 * i), line, fill="black", font=font)
    img.save(path)
    return path


def make_scanned_pdf(path: Path, lines: list[str]) -> Path:
    """An image-only PDF (no text layer)."""
    from PIL import Image

    png = make_png(path.with_suffix(".scan.png"), lines)
    Image.open(png).convert("RGB").save(path, "PDF")
    return path


def make_xlsx(path: Path, sheets: dict[str, list[list]]) -> Path:
    import openpyxl

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
    wb.save(path)
    return path


def make_docx(path: Path, blocks: list[tuple]) -> Path:
    import docx

    d = docx.Document()
    for kind, value in blocks:
        if kind == "h1":
            d.add_heading(value, level=1)
        elif kind == "h2":
            d.add_heading(value, level=2)
        elif kind == "p":
            d.add_paragraph(value)
        elif kind == "table":
            t = d.add_table(rows=len(value), cols=len(value[0]))
            for r, row in enumerate(value):
                for c, cell in enumerate(row):
                    t.cell(r, c).text = str(cell)
    d.save(path)
    return path


def make_pptx(path: Path, slides: list[tuple[str, str]]) -> Path:
    import pptx

    prs = pptx.Presentation()
    for title, body in slides:
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = title
        s.placeholders[1].text = body
    prs.save(path)
    return path


class FakeModel:
    """Stand-in for a model call: records every call; replies with `reply`, or raises when `fail`."""

    def __init__(self, reply: str = "FAKE MARKDOWN", fail: bool = False):
        self.reply = reply
        self.fail = fail
        self.calls: list[tuple] = []

    def __call__(self, cli, model, prompt_or_image, *args, **kw):
        self.calls.append((cli, model, prompt_or_image, args, kw))
        if self.fail:
            from bron.kb.models import ModelError

            raise ModelError("the model is not available")
        return self.reply


class FakeEmbedder:
    """Deterministic stand-in for the meaning model: hashes words into 384 dims, normalised. Ranks exact-word overlap."""

    model = "fake-embedder"

    def __init__(self):
        self.calls = 0
        self.seen: list[str] = []

    def embed(self, texts):
        import re
        import unicodedata
        import zlib

        import numpy as np

        self.calls += 1
        self.seen.extend(texts)
        out = np.zeros((len(texts), 384), dtype=np.float32)
        for i, text in enumerate(texts):
            folded = "".join(c for c in unicodedata.normalize("NFKD", text) if not unicodedata.combining(c)).casefold()
            for word in re.findall(r"\w+", folded):
                out[i, zlib.crc32(word.encode()) % 384] += 1.0
        norms = np.linalg.norm(out, axis=1, keepdims=True)
        return out / np.where(norms == 0, 1, norms)


fake_embed = FakeEmbedder()


SCAN_TEXT = ("Termo de investimento: a Acme recebeu R$ 1.500.000,00 em 21 de janeiro de 2025, "
             "conforme a cláusula 4.2 do acordo de acionistas.")


class FakeOcr:
    """Stand-in for Mac text recognition: records every image it was given; replies with `text`."""

    def __init__(self, text: str = SCAN_TEXT, confidence: float = 0.95):
        self.text = text
        self.confidence = confidence
        self.images: list[Path] = []

    def __call__(self, image):
        self.images.append(Path(image))
        assert Path(image).is_file()
        return self.text, self.confidence


def hook_reads(monkeypatch, before) -> None:
    """Run before(<document name>) each time Bron starts reading a document (raise in it to stop the reader there)."""
    from bron.kb import ingest

    real = ingest._read

    def hooked(item, ocr, work):
        before(item.name)
        return real(item, ocr, work)

    monkeypatch.setattr(ingest, "_read", hooked)


def write_page(vault, rel: str, body: str = "", **meta) -> Path:
    """A wiki page under Knowledge/ (rel like "Organisations/Acme Ltda.md") with these properties."""
    from bron import frontmatter as fm

    path = vault.knowledge_dir / rel
    fm.write(path, fm.Document(dict(meta), body))
    return path


def later(path: Path, seconds: float = 5.0) -> None:
    """Move a file's modification time forward, so a quick edit in a test always counts as a change."""
    st = path.stat()
    os.utime(path, (st.st_atime + seconds, st.st_mtime + seconds))


def stored_doc(vault, name: str, pages: list[str], *, folder: str = "", index_it: bool = False):
    """A document Bron has read (in the store; in the search index too with index_it), without any reader."""
    from bron.kb import index, store
    from bron.kb.models import labels_from_names
    from bron.kb.passages import split

    identity = f"file:/docs/{name}"
    doc = store.Doc(store.doc_id_for(identity), identity, "file", f"/docs/{name}", name, f"/docs/{name}",
                    labels=labels_from_names(name, folder), read_at="2026-10-04", pages=len(pages))
    passages = split([(i + 1, t) for i, t in enumerate(pages)], store.effective_labels(doc))
    store.save(vault, doc, pages, passages)
    if index_it:
        index.put(vault, doc, passages, fake_embed)
    return store.load(vault, doc.doc_id)


class FakeTicketRunner:
    """Stand-in for runner.run_ticket: records the tickets it runs, calls write(vault, ticket) first (to write pages,
    say), then settles the ticket like a real run: in-review with a summary, or `status` with `message`."""

    def __init__(self, status: str = "in-review", message: str = "", write=None):
        self.status, self.message, self.write = status, message, write
        self.calls: list[str] = []

    def __call__(self, vault, ticket_id, **kw):
        from bron.runner import RunOutcome
        from bron.tickets import editing, find_ticket, load_ticket, set_result, set_status

        self.calls.append(ticket_id)
        if self.write is not None:
            self.write(vault, load_ticket(find_ticket(vault, ticket_id)))
        with editing(vault, ticket_id) as ticket:
            if self.status == "in-review":
                set_result(ticket, "Learned: three leases. Pages: 3 documents, 1 organisation.", ticket.assignee)
            else:
                set_status(ticket, self.status, "runner", self.message or "stand-in")
        return RunOutcome(ticket_id, self.status, "claude", self.message or f"{ticket_id} is now {self.status}")
