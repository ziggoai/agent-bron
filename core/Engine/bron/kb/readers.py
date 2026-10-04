"""One reader per file type. Scans and photos go through the Mac's text recognition, page by page."""
from __future__ import annotations

import csv
import datetime as dt
import hashlib
import html as htmllib
import io
import re
import subprocess
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable

from .store import KbError

SHEET_ROWS = 50
WORD_PART_CHARS = 3000
MIN_TEXT_CHARS = 40
THIN_TEXT_CHARS = 200
BAD_SHARE = 0.05

PDF_EXT = {".pdf"}
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".heic", ".heif", ".tif", ".tiff", ".bmp", ".gif", ".webp"}
SHEET_EXT = {".xlsx", ".xlsm", ".xls", ".csv"}
TEXT_EXT = {".txt", ".md"}
HTML_EXT = {".html", ".htm"}


@dataclass
class Page:
    number: int
    text: str
    ocr: bool = False
    image: Path | None = None
    confidence: float = 1.0


@dataclass
class Read:
    pages: list[Page] = field(default_factory=list)
    kind: str = "text"


Ocr = Callable[[Path], "tuple[str, float]"]


OLE2_MAGIC = bytes.fromhex("D0CF11E0A1B11AE1")
MODERN_OFFICE = {".xlsx", ".xlsm", ".docx", ".pptx"}


def _damaged(kind: str) -> KbError:
    article = "an" if kind[:1].lower() in "aeiou" else "a"
    return KbError(f"This file is damaged or isn't really {article} {kind}.")


def _check_not_encrypted(path: Path) -> None:
    """Password-protected modern Office files are OLE2 containers instead of zip files."""
    with open(path, "rb") as fh:
        if fh.read(8) == OLE2_MAGIC:
            raise KbError("This file is password-protected.")


# ---- Mac text recognition ----

def ocr_page(image: Path) -> tuple[str, float]:
    """Text of one image (lines joined with newlines) and the mean confidence."""
    from ocrmac import ocrmac

    found = ocrmac.OCR(str(image), language_preference=["pt-BR", "en-US"], recognition_level="accurate").recognize()
    if not found:
        return "", 0.0
    return "\n".join(t for t, _, _ in found), sum(c for _, c, _ in found) / len(found)


def _run_ocr(ocr: Ocr, image: Path) -> tuple[str, float]:
    try:
        return ocr(image)
    except KbError:
        raise
    except Exception as exc:
        raise KbError("Bron couldn't read the text in this scan.") from exc


# ---- dispatch ----

def read(path: Path, *, ocr: Ocr | None = None, work: Path) -> Read:
    path = Path(path)
    ext = path.suffix.lower()
    ocr = ocr or ocr_page
    work = Path(work)
    work.mkdir(parents=True, exist_ok=True)
    if ext in MODERN_OFFICE:
        _check_not_encrypted(path)
    if ext in PDF_EXT:
        return _read_pdf(path, ocr, work)
    if ext in IMAGE_EXT:
        return _read_image(path, ocr, work)
    if ext in SHEET_EXT:
        return _read_sheet(path, ext)
    if ext == ".docx":
        return _read_docx(path)
    if ext == ".pptx":
        return _read_pptx(path)
    if ext in TEXT_EXT:
        return Read(_one_page(decode(path.read_bytes())), "text")
    if ext in HTML_EXT:
        return read_html(decode(path.read_bytes()), "")
    raise KbError(f"Bron can't read {ext or 'these'} files yet.")


def decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", "replace")


def _one_page(text: str) -> list[Page]:
    return [Page(1, text)] if text.strip() else []


# ---- PDF ----

def _needs_ocr(text: str) -> bool:
    stripped = text.strip()
    if len(stripped) < MIN_TEXT_CHARS:
        return True
    bad = sum(1 for ch in stripped if ch == "�" or (ord(ch) < 32 and ch not in "\n\r\t"))
    if bad / len(stripped) > BAD_SHARE:
        return True
    if "(cid:" in stripped:
        return True
    # dot leaders, rules and signature lines are layout, not garbage
    solid = list(re.sub(r"([^\w\s]|_)\1{2,}", "", re.sub(r"\s+", "", stripped)))
    if not solid:
        return True
    if sum(1 for ch in solid if "" <= ch <= "") / len(solid) >= 0.10:
        return True
    return sum(1 for ch in solid if ch.isalnum()) / len(solid) < 0.5


def _covered_by_image(page) -> bool:
    """True when one image covers at least half of the page (a scan under a thin text header)."""
    import pypdfium2.raw as pdfium_raw

    width, height = page.get_size()
    area = width * height
    if area <= 0:
        return False
    for obj in page.get_objects(filter=[pdfium_raw.FPDF_PAGEOBJ_IMAGE]):
        left, bottom, right, top = obj.get_bounds()
        if (right - left) * (top - bottom) >= 0.5 * area:
            return True
    return False


def _read_pdf(path: Path, ocr: Ocr, work: Path) -> Read:
    import pypdfium2 as pdfium

    try:
        doc = pdfium.PdfDocument(path)
    except pdfium.PdfiumError as exc:
        if "password" in str(exc).lower():
            raise KbError("This PDF is password-protected.") from exc
        raise _damaged("PDF") from exc
    except Exception as exc:
        raise _damaged("PDF") from exc
    pages: list[Page] = []
    try:
        for index in range(len(doc)):
            number = index + 1
            page = doc[index]
            try:
                text = page.get_textpage().get_text_range()
                if not _needs_ocr(text) and not (len(text.strip()) < THIN_TEXT_CHARS and _covered_by_image(page)):
                    pages.append(Page(number, text))
                    continue
                image = work / f"p{number}.png"
                page.render(scale=2).to_pil().save(image)
            finally:
                page.close()
            recognised, confidence = _run_ocr(ocr, image)
            pages.append(Page(number, recognised, ocr=True, image=image, confidence=confidence))
    except KbError:
        raise
    except pdfium.PdfiumError as exc:
        raise _damaged("PDF") from exc
    finally:
        doc.close()
    return Read(pages, "pdf")


# ---- images ----

def _read_image(path: Path, ocr: Ocr, work: Path) -> Read:
    image = path
    if path.suffix.lower() in {".heic", ".heif"}:
        tag = hashlib.sha256(str(path.resolve()).encode("utf-8")).hexdigest()[:8]
        image = work / f"{path.stem}-{tag}.png"
        done = subprocess.run(["sips", "-s", "format", "png", str(path), "--out", str(image)],
                              capture_output=True, text=True)
        if done.returncode != 0 or not image.exists():
            raise _damaged("image")
    else:
        try:
            from PIL import Image

            with Image.open(path) as im:
                im.verify()
        except ImportError:
            pass
        except Exception as exc:
            raise _damaged("image") from exc
    text, confidence = _run_ocr(ocr, image)
    return Read([Page(1, text, ocr=True, image=image, confidence=confidence)], "image")


# ---- sheets ----

def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer() and abs(value) < 1e18:
        return str(int(value))
    if isinstance(value, dt.datetime):
        return value.date().isoformat() if value.time() == dt.time(0) else value.isoformat(sep=" ")
    if isinstance(value, (dt.date, dt.time)):
        return value.isoformat()
    return str(value).replace("\r\n", " ").replace("\n", " ").replace("|", "\\|").strip()


def _md_table(rows: list[list[str]]) -> str:
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    lines = ["| " + " | ".join(rows[0]) + " |", "|" + " --- |" * width]
    lines += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return "\n".join(lines)


def _sheet_pages(name: str, raw_rows: list[list], start: int) -> list[Page]:
    rows = [[_cell(v) for v in r] for r in raw_rows]
    rows = [r for r in rows if any(c for c in r)]
    if not rows:
        return []
    header, body = rows[0], rows[1:]
    chunks = [body[i:i + SHEET_ROWS] for i in range(0, len(body), SHEET_ROWS)] or [[]]
    return [Page(start + i, f"Sheet: {name}\n\n" + _md_table([header] + chunk)) for i, chunk in enumerate(chunks)]


def _read_sheet(path: Path, ext: str) -> Read:
    sheets: list[tuple[str, list[list]]] = []
    if ext == ".csv":
        text = decode(path.read_bytes())
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        sheets.append((path.stem, [list(r) for r in csv.reader(io.StringIO(text, newline=""), dialect)]))
    else:
        from python_calamine import CalamineWorkbook

        try:
            book = CalamineWorkbook.from_path(str(path))
            for name in book.sheet_names:
                sheets.append((name, book.get_sheet_by_name(name).to_python()))
        except Exception as exc:
            raise _damaged("spreadsheet") from exc
    pages: list[Page] = []
    for name, rows in sheets:
        pages += _sheet_pages(name, rows, len(pages) + 1)
    return Read(pages, "sheet")


# ---- Word ----

def split_parts(blocks: list[str]) -> list[Page]:
    pages: list[Page] = []
    current: list[str] = []
    size = 0
    for block in blocks:
        if current and size + len(block) > WORD_PART_CHARS:
            pages.append(Page(len(pages) + 1, "\n\n".join(current)))
            current, size = [], 0
        current.append(block)
        size += len(block) + 2
    if current:
        pages.append(Page(len(pages) + 1, "\n\n".join(current)))
    return pages


def _read_docx(path: Path) -> Read:
    try:
        import docx
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        doc = docx.Document(str(path))
        blocks: list[str] = []
        for child in doc.element.body.iterchildren():
            tag = child.tag.rsplit("}", 1)[-1]
            if tag == "p":
                para = Paragraph(child, doc)
                text = para.text.strip()
                if not text:
                    continue
                style = (para.style.name if para.style is not None else "") or ""
                match = re.match(r"Heading (\d)", style)
                if match:
                    text = "#" * min(int(match.group(1)), 6) + " " + text
                elif style == "Title":
                    text = "# " + text
                blocks.append(text)
            elif tag == "tbl":
                rows = [[_cell(c.text) for c in row.cells] for row in Table(child, doc).rows]
                if rows:
                    blocks.append(_md_table(rows))
    except Exception as exc:
        raise _damaged("Word document") from exc
    return Read(split_parts(blocks), "word")


# ---- PowerPoint ----

def _shape_texts(shapes, skip) -> list[str]:
    out: list[str] = []
    for shape in shapes:
        if shape is skip:
            continue
        if getattr(shape, "shapes", None) is not None and shape.shape_type == 6:
            out += _shape_texts(shape.shapes, skip)
        elif getattr(shape, "has_table", False) and shape.has_table:
            out.append(_md_table([[_cell(c.text) for c in row.cells] for row in shape.table.rows]))
        elif getattr(shape, "has_text_frame", False) and shape.has_text_frame:
            text = shape.text_frame.text.strip()
            if text:
                out.append(text)
    return out


def _read_pptx(path: Path) -> Read:
    try:
        import pptx

        prs = pptx.Presentation(str(path))
        pages: list[Page] = []
        for number, slide in enumerate(prs.slides, start=1):
            title_shape = slide.shapes.title
            parts: list[str] = []
            if title_shape is not None and title_shape.text_frame.text.strip():
                parts.append("# " + title_shape.text_frame.text.strip())
            parts += _shape_texts(slide.shapes, title_shape)
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    parts.append("Notes: " + notes)
            pages.append(Page(number, "\n\n".join(parts)))
    except Exception as exc:
        raise _damaged("PowerPoint file") from exc
    return Read(pages, "slides")


# ---- web pages ----

def _strip_tags(html: str) -> str:
    html = re.sub(r"(?is)<(script|style|noscript|head)\b.*?</\1>", " ", html)
    html = re.sub(r"(?i)</(p|div|h\d|li|tr|br)\s*>|<br\s*/?>", "\n", html)
    text = htmllib.unescape(re.sub(r"<[^>]+>", " ", html))
    lines = (re.sub(r"[ \t]+", " ", ln).strip() for ln in text.splitlines())
    return "\n".join(ln for ln in lines if ln)


def read_html(html: str, url: str) -> Read:
    text = ""
    try:
        import trafilatura

        text = trafilatura.extract(html, url=url or None, include_tables=True) or ""
    except Exception:
        text = ""
    if not text.strip():
        text = _strip_tags(html)
    return Read(_one_page(text), "web")
